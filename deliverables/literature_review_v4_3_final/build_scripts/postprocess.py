"""Post-process the pandoc DOCX: table formatting, widths, captions, bibliography, footer page numbers."""
import sys, re, copy
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, dst = sys.argv[1], sys.argv[2]
d = Document(src)
sec = d.sections[0]
TEXT_W = sec.page_width - sec.left_margin - sec.right_margin  # EMU
TW_DXA = int(TEXT_W / 635)  # EMU -> twips

def set_cell_props(cell, shade=None):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side in ("top", "bottom"):
        e = OxmlElement("w:%s" % side); e.set(qn("w:w"), "28"); e.set(qn("w:type"), "dxa"); mar.append(e)
    for side in ("left", "right"):
        e = OxmlElement("w:%s" % side); e.set(qn("w:w"), "57"); e.set(qn("w:type"), "dxa"); mar.append(e)
    tcPr.append(mar)
    if shade:
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), shade)
        tcPr.append(shd)

def borders(tbl):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for side, sz, col in (("top", 10, "1F2A44"), ("bottom", 10, "1F2A44"), ("insideH", 3, "A0A8B8"), ("left", 0, "auto"), ("right", 0, "auto"), ("insideV", 3, "D0D4DC")):
        e = OxmlElement("w:%s" % side)
        e.set(qn("w:val"), "single" if sz else "nil"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), col)
        b.append(e)
    tblPr.append(b)

def col_widths(tbl):
    rows = tbl.rows
    n = len(rows[0].cells)
    lens = [0.0] * n
    for r in rows:
        for j, c in enumerate(r.cells[:n]):
            t = c.text.strip()
            words = t.split()
            longest = max((len(w) for w in words), default=1)
            lens[j] = max(lens[j], min(len(t), 60) * 0.65 + longest * 0.9, 4)
    longest = [0] * n
    for r in rows:
        for j, c in enumerate(r.cells[:n]):
            longest[j] = max([longest[j]] + [len(x) for x in c.text.split()])
    mins = [lw * 78 + 150 for lw in longest]
    s = sum(lens)
    w = [max(int(TW_DXA * l / s), m) for l, m in zip(lens, mins)]
    excess = sum(w) - TW_DXA
    if excess > 0:
        slack = [max(x - m, 0) for x, m in zip(w, mins)]
        tot = sum(slack) or 1
        w = [int(x - excess * sl / tot) for x, sl in zip(w, slack)]
    return w

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
for ti, tbl in enumerate(d.tables):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(tbl)
    tblPr = tbl._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for old in tblPr.findall(qn(tag)):
            tblPr.remove(old)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(TW_DXA)); tw.set(qn("w:type"), "dxa"); tblPr.append(tw)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    widths = col_widths(tbl)
    if ti == 0:  # notation table: symbols are OMML (no plain text), so fix proportions
        widths = [int(TW_DXA * f) for f in (0.27, 0.58, 0.15)]
    # shrink math in tables
    for mr in tbl._tbl.iter("{%s}r" % M_NS):
        rpr = mr.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr"); mr.insert(0, rpr) if mr.find("{%s}rPr" % M_NS) is None else mr.insert(1, rpr)
        for tag in ("w:sz", "w:szCs"):
            e = OxmlElement(tag); e.set(qn("w:val"), "17"); rpr.append(e)
    grid = tbl._tbl.tblGrid
    for gc, wv in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(wv))
    ncols = len(widths)
    small = ncols >= 6
    for i, row in enumerate(tbl.rows):
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit"); trPr.append(cant)
        if i == 0:
            h = OxmlElement("w:tblHeader"); trPr.append(h)
        for j, cell in enumerate(row.cells):
            cell.width = Emu(widths[min(j, ncols - 1)] * 635)
            set_cell_props(cell, "E8ECF3" if i == 0 else None)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(7.5 if small else 8.0)
                    if i == 0: r.font.bold = True

# Captions: paragraphs starting with "Table N." -> keep with next, caption formatting
for p in d.paragraphs:
    t = p.text.strip()
    if re.match(r"^Table [A-Z]?\d+\.", t):
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.size = Pt(9)
    if p.style.name in ("Image Caption", "Captioned Figure", "Figure"):
        p.paragraph_format.keep_with_next = False
    if p.style.name == "Bibliography":
        pf = p.paragraph_format
        pf.left_indent = Cm(0.9); pf.first_line_indent = Cm(-0.9); pf.space_after = Pt(2.5)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.size = Pt(8.5)
    if p.style.name == "Abstract":
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.right_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(5)
        for r in p.runs: r.font.size = Pt(9.5)
    if p.style.name == "Block Text":
        p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.right_indent = Cm(0.6)
        for r in p.runs: r.font.size = Pt(9)
    # images: keep figure paragraph with caption
    if p._p.findall(".//" + qn("w:drawing")):
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Footer page numbers
footer = sec.footer
fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
def fld(par, instr):
    r = par.add_run(); f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r._r.append(f1)
    r2 = par.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r2._r.append(it)
    r3 = par.add_run(); f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end"); r3._r.append(f2)
run = fp.add_run("V4.3 literature-review manuscript — page "); run.font.size = Pt(8)
fld(fp, "PAGE")
for r in fp.runs: r.font.size = Pt(8)

# core properties
cp = d.core_properties
cp.title = "From Detection to Mechanism Diagnosis: An Observability- and Payload-Aware Critical Review of RF Interference Recognition in GEO Bent-Pipe Satellite Systems"
cp.subject = "Literature review manuscript V4.3"
cp.author = "Burak Tos"
cp.keywords = "GEO satellite; bent-pipe transponder; RF interference; observability; open-set recognition"
cp.comments = "V4.3 literature-review manuscript; no new experimental results."
d.save(dst)
print("tables", len(d.tables))
